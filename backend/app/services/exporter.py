from io import BytesIO

from docx import Document

from app.models import LibraryEntry


def library_to_markdown(entries: list[LibraryEntry]) -> str:
    parts: list[str] = ["# AI Research Paper Summaries"]
    for entry in entries:
        paper = entry.paper
        summary = entry.summary
        parts.extend(
            [
                "",
                f"## {paper.title}",
                f"- arXiv: {paper.id}",
                f"- Authors: {', '.join(paper.authors)}",
                f"- Categories: {', '.join(paper.categories)}",
                "",
                "### Summary",
                *[f"- {item}" for item in summary.summary_bullets],
                "",
                "### Contributions",
                *[f"- {item}" for item in summary.contributions],
                "",
                "### Limitations",
                *[f"- {item}" for item in summary.limitations],
                "",
                "### Follow-up Papers",
                *[f"- {item}" for item in summary.follow_up_papers],
            ]
        )
    return "\n".join(parts)


def library_to_docx(entries: list[LibraryEntry]) -> BytesIO:
    doc = Document()
    doc.add_heading("AI Research Paper Summaries", 0)
    for entry in entries:
        paper = entry.paper
        summary = entry.summary
        doc.add_heading(paper.title, level=1)
        doc.add_paragraph(f"arXiv: {paper.id}")
        doc.add_paragraph(f"Authors: {', '.join(paper.authors)}")
        doc.add_paragraph(f"Categories: {', '.join(paper.categories)}")

        for heading, items in (
            ("Summary", summary.summary_bullets),
            ("Contributions", summary.contributions),
            ("Limitations", summary.limitations),
            ("Follow-up Papers", summary.follow_up_papers),
        ):
            doc.add_heading(heading, level=2)
            for item in items:
                doc.add_paragraph(item, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
